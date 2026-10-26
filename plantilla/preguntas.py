# Genera la plantilla Excel de ejemplo y la semilla de preguntas para Code.gs
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

PREGUNTAS = [
 ("Un sistema de información incluye personas, procedimientos, datos y tecnología, no solo computadoras.", "V"),
 ("Dato e información son sinónimos en la gestión empresarial.", "F"),
 ("La información tiene valor cuando es oportuna, precisa y relevante para quien decide.", "V"),
 ("Un ERP integra en una única base de datos los procesos de distintas áreas de la empresa.", "V"),
 ("Un ERP se implementa siempre sin necesidad de revisar los procesos de negocio de la empresa.", "F"),
 ("La memoria RAM conserva su contenido cuando se apaga el equipo.", "F"),
 ("Un disco SSD no tiene partes mecánicas móviles.", "V"),
 ("La computación en la nube permite pagar por uso en lugar de invertir en infraestructura propia.", "V"),
 ("En el modelo SaaS, el cliente es responsable de administrar los servidores donde corre la aplicación.", "F"),
 ("IaaS provee infraestructura (servidores, almacenamiento, red) como servicio.", "V"),
 ("El sistema operativo es un software de aplicación.", "F"),
 ("El software de código abierto nunca tiene costos asociados de implementación ni de soporte.", "F"),
 ("Un algoritmo es una secuencia finita de pasos para resolver un problema.", "V"),
 ("RFID permite identificar objetos sin línea de visión directa, a diferencia del código de barras.", "V"),
 ("La mayor parte del tráfico internacional de Internet viaja por satélites.", "F"),
 ("TCP/IP es el conjunto de protocolos sobre el que funciona Internet.", "V"),
 ("Una red LAN cubre áreas geográficas extensas, como países o continentes.", "F"),
 ("La Business Intelligence transforma datos en información para la toma de decisiones.", "V"),
 ("Un data warehouse está optimizado principalmente para registrar transacciones operativas en tiempo real.", "F"),
 ("La minería de datos busca patrones y relaciones no evidentes en grandes volúmenes de datos.", "V"),
 ("Un tablero de control debería mostrar todos los indicadores disponibles, sin priorizar.", "F"),
 ("La inteligencia artificial generativa puede producir respuestas incorrectas con apariencia convincente.", "V"),
 ("COBIT es un marco de referencia para el gobierno y la gestión de las TI.", "V"),
 ("La gobernanza de IT es responsabilidad exclusiva del área de sistemas.", "F"),
 ("Alinear la estrategia de IT con la estrategia del negocio es un objetivo central de la gobernanza de IT.", "V"),
 ("El CHAOS Report del Standish Group analiza el éxito y el fracaso de proyectos de software.", "V"),
 ("Según el CHAOS Report, la mayoría de los proyectos de software terminan en el tiempo, costo y alcance previstos.", "F"),
 ("Las metodologías ágiles entregan el producto en incrementos cortos e iterativos.", "V"),
 ("En una metodología en cascada, cada etapa comienza cuando finaliza la anterior.", "V"),
 ("Scrum es una metodología en cascada.", "F"),
 ("El involucramiento de los usuarios es un factor clave de éxito en los proyectos de sistemas.", "V"),
 ("Desarrollar software a medida suele ser más rápido y barato que adquirir un paquete estándar.", "F"),
 ("Al tercerizar un servicio de IT, la empresa deja de necesitar controlar al proveedor.", "F"),
 ("Un acuerdo de nivel de servicio (SLA) define parámetros medibles de calidad que el proveedor se compromete a cumplir.", "V"),
 ("El costo total de propiedad (TCO) considera solo el precio de compra del equipamiento.", "F"),
 ("En una contratación de tecnología conviene definir los criterios de evaluación antes de recibir las propuestas.", "V"),
 ("Confidencialidad, integridad y disponibilidad son los pilares de la seguridad de la información.", "V"),
 ("Un firewall alcanza por sí solo para proteger a una organización de todos los ataques informáticos.", "F"),
 ("El phishing busca engañar a las personas para que entreguen credenciales o información sensible.", "V"),
 ("El ransomware cifra los datos de la víctima y exige un pago para liberarlos.", "V"),
 ("Tener backups probados periódicamente es una medida clave frente a un ataque de ransomware.", "V"),
 ("La autenticación de dos factores combina, por ejemplo, algo que sabés con algo que tenés.", "V"),
 ("Un plan de continuidad del negocio solo es necesario en empresas de más de mil empleados.", "F"),
 ("El e-business abarca únicamente la venta de productos por Internet.", "F"),
 ("El modelo B2B se refiere a transacciones entre empresas.", "V"),
 ("Las plataformas digitales pueden convertir en commodity a las empresas que dependen de ellas para llegar a sus clientes.", "V"),
 ("La Industria 4.0 se caracteriza por integrar sistemas físicos y digitales, como sensores conectados e IoT.", "V"),
 ("Internet de las Cosas (IoT) se refiere exclusivamente a la conexión de computadoras personales a Internet.", "F"),
 ("El diagrama de un circuito administrativo permite detectar controles internos y puntos de mejora del proceso.", "V"),
 ("Digitalizar un proceso ineficiente lo vuelve eficiente automáticamente.", "F"),
]
assert len(PREGUNTAS) == 50

def xlsx(path):
    wb = Workbook(); ws = wb.active; ws.title = "Preguntas"
    ws.append(["N", "Pregunta", "Respuesta"])
    for i, (q, a) in enumerate(PREGUNTAS, 1):
        ws.append([i, q, a])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="16202A")
    ws.column_dimensions["A"].width = 6; ws.column_dimensions["B"].width = 100; ws.column_dimensions["C"].width = 12
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top"); row[2].alignment = Alignment(horizontal="center")
    dv = DataValidation(type="list", formula1='"V,F"', allow_blank=False); ws.add_data_validation(dv); dv.add("C2:C500")
    ws.freeze_panes = "A2"
    ins = wb.create_sheet("Instrucciones")
    for line in ["Cómo completar la plantilla",
                 "1. Una pregunta por fila en la hoja Preguntas (columna Pregunta).",
                 "2. En Respuesta poné V (verdadero) o F (falso). También se aceptan Verdadero / Falso.",
                 "3. La columna N es opcional: el orden de las filas define el número de pregunta.",
                 "4. Subí el archivo desde la sección Preguntas del panel de administración.",
                 "Estas 50 preguntas son de ejemplo: reemplazalas por las tuyas antes de tomar el examen."]:
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=13); ins.column_dimensions["A"].width = 100
    wb.save(path)

if __name__ == "__main__":
    xlsx(sys.argv[1])
    print(json.dumps([[q, a] for q, a in PREGUNTAS], ensure_ascii=False))
